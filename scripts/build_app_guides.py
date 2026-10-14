from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'docs'
OUT.mkdir(exist_ok=True)
APP = 'https://lawrence-resume-studio.spiffydace4.chatgpt.site'
DATE = 'October 14 2026'

def para(doc, text, style=None):
    p = doc.add_paragraph(text, style)
    return p

def h(doc, text, level=1):
    return doc.add_heading(text, level)

def bullets(doc, items):
    for item in items:
        para(doc, item, 'List Bullet')

def steps(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.22)
        p.paragraph_format.first_line_indent = Inches(-.22)
        p.add_run(str(i) + '. ').bold = True
        p.add_run(item)

def code(doc, text):
    p = doc.add_paragraph(style='Code')
    p.add_run(text)
    return p

def link(doc, label, url, prefix=''):
    p = doc.add_paragraph(prefix)
    rel = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); props.append(u)
    r.append(props); t = OxmlElement('w:t'); t.text = label; r.append(t); hl.append(r); p._p.append(hl)
    return p

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns,widths): column.width = Inches(width)
    for cell, title, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width); cell.text = title
    head = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); head.append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width); cell.text = text
    for ri, row in enumerate(t.rows):
        prop = row._tr.get_or_add_trPr(); no_split = OxmlElement('w:cantSplit'); prop.append(no_split)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, value in [('top','85'),('left','110'),('bottom','85'),('right','110')]:
                x = OxmlElement('w:'+side); x.set(qn('w:w'),value); x.set(qn('w:type'),'dxa'); margins.append(x)
            cp.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                x = OxmlElement('w:'+side); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D0D0D0'); borders.append(x)
            cp.append(borders)
            if ri == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'F2F2F2'); cp.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.06
                if headers[-1]=='Status' and ci==len(headers)-1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Arial'; run.font.size = Pt(10.5); run.font.color.rgb = RGBColor(0,0,0)
                    run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def base(kind, opening):
    d = Document(); sec = d.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7)
    sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    sec.header_distance = Inches(.3); sec.footer_distance = Inches(.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
        s = d.styles[name]; s.font.name = 'Arial'; s.font.color.rgb = RGBColor(0,0,0)
        s.paragraph_format.widow_control = True
        if name.startswith('Heading') or name in ['Title','Subtitle']: s.paragraph_format.keep_with_next = True
    d.styles['Normal'].font.size = Pt(11)
    d.styles['Normal'].paragraph_format.line_spacing = 1.08
    d.styles['Normal'].paragraph_format.space_after = Pt(7)
    d.styles['Title'].font.size = Pt(25); d.styles['Title'].font.bold = True
    d.styles['Title'].paragraph_format.space_after = Pt(5)
    d.styles['Subtitle'].font.size = Pt(12); d.styles['Subtitle'].paragraph_format.space_after = Pt(14)
    d.styles['Heading 1'].font.size = Pt(16); d.styles['Heading 1'].font.bold = True
    d.styles['Heading 1'].paragraph_format.space_before = Pt(13); d.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    d.styles['Heading 2'].font.size = Pt(12); d.styles['Heading 2'].font.bold = True
    d.styles['Heading 2'].paragraph_format.space_before = Pt(10); d.styles['Heading 2'].paragraph_format.space_after = Pt(5)
    d.styles['List Bullet'].font.size = Pt(11)
    d.styles['List Bullet'].paragraph_format.space_after = Pt(4)
    cs = d.styles.add_style('Code',1); cs.font.name = 'Courier New'; cs.font.size = Pt(9.5); cs.font.color.rgb = RGBColor(0,0,0)
    cs.paragraph_format.line_spacing = 1.04; cs.paragraph_format.space_after = Pt(8)
    cs.paragraph_format.keep_together = True
    for s in d.styles:
        if s.type == 1:
            pr = s.element.find(qn('w:pPr'))
            if pr is not None:
                for child in list(pr):
                    if child.tag == qn('w:pBdr'): pr.remove(child)
    header = sec.header.paragraphs[0]
    header.add_run('Le Resume Studio  |  '+kind).font.size = Pt(9)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('© 2026 Marabou Lane LLC   |   '); r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    d.core_properties.title = 'Le Resume Studio '+kind
    d.core_properties.author = 'Marabou Lane LLC'; d.core_properties.subject = opening[:255]
    para(d, 'Le Resume Studio', 'Title'); para(d, kind, 'Subtitle')
    para(d, DATE+'  |  Published app version 21')
    para(d, opening)
    return d

def new_page(d, title):
    p=h(d,title); p.paragraph_format.page_break_before=True

def save(d, name):
    path = OUT / (name+'.docx'); d.save(path)
    print(path)

def tutorial():
    d = base('User Tutorial', 'Use this guide to turn your own work history into a tailored resume and matching cover letter, review the results, and save both files. A new visitor starts with a blank master resume. Returning visitors keep the master and application versions saved in their browser.')
    link(d, 'Open Le Resume Studio', APP)
    h(d, '1 Add your master resume')
    para(d, 'Open Master resume or choose Add resume. Upload a Word (.docx), text-based PDF, UTF-8 .txt or rich-text .rtf resume. Master resume allows file upload only; there are no manual editing fields. Files must be nonempty and no larger than 10 MB; PDFs must have 30 pages or fewer. Scans need a text copy. Uploads are read in your browser without sending resume contents to a server.')
    table(d, ['Section','Optional information the app can use'], [
        ('Contact details','Name, email, phone and LinkedIn URL are optional for upload. Review contact information before applying.'),
        ('Summary and competencies','A summary and skills are optional. Include them when available; no specific headings are required.'),
        ('Experience','Job titles, companies, dates and accomplishments are optional for upload. Supplied facts are used without inventing missing details.'),
        ('Skills and education','Skills and education are optional. Missing school names, qualifications and dates do not block upload.')
    ], [1.45,5.55])
    para(d, 'Any readable nonempty upload saves automatically, without required sections. Review the version-one-style sections for contact details, summary, competencies, experience, technical skills and education. Each detected role has its own expanded Experience section showing company, years and bullets. View uploaded resume text shows the original extraction. All source details are read-only. Final accomplishments stay with their role; job headings retain the uploaded role, company and year order. To correct the source, edit your original file and upload it again. Each successful upload requires a unique filename in this browser master workspace, ignoring capitalization. Rename revised files before uploading. Clear Master Resume also clears filename history. Failed uploads preserve the prior master.')
    para(d, 'Choose Clear Master Resume to immediately delete every master subsection, imported file data, extracted text and filename history from this app. No browser confirmation popup is needed. The dialog shows No master resume saved; the current workspace clears and pending imports are canceled. The master stays blank after reopening. Failed storage deletion shows an error and allows retry. Your source file on your computer is unchanged; the app retains extracted data rather than the original file bytes. Saved applications are removed separately with Clear saved versions. Word and PDF uploads retain supported typography and structure; complex layouts can change. TXT and RTF use a professional single-column template.')

    new_page(d, 'Protect your private workspace')
    para(d, 'App links do not contain applicant data. Uploads, tailoring and exports run in your browser; the app does not send resumes to an applicant database or other users. A fresh browser starts blank. On a shared browser profile, enable password protection before leaving saved data on the device.')
    steps(d, ['Choose Privacy and password in the navigation.', 'Enter and confirm a unique passphrase of at least 12 characters, then choose Protect saved data. Existing master and saved versions are encrypted; their old plaintext records are removed after encryption succeeds.', 'Before handing the device to another person, choose Privacy and password, then Lock private workspace. This blanks the workspace without deleting the encrypted master or saved versions.', 'After reloading or reopening, enter your passphrase through Privacy and password to unlock the saved data. Sections 1 and 2 remain blank until you use the master or reopen a saved version.'])
    para(d, 'Keep the passphrase safely: there is no recovery service, server account or cross-device synchronization. Without password protection, saved data belongs to the browser profile and can be read by another person using that profile. Protection applies to saved data; another person can see an unlocked tab or a separately downloaded file. Close download tabs and protect files on the device separately.')
    para(d, 'Open download page uses a private, five-minute browser-session handoff. The URL contains only a random identifier and cannot retrieve the document in another browser session. Popups or session storage may be blocked; use the direct save option if the new tab cannot open. Sharing a file through your device’s Share option is an explicit action you control.')

    new_page(d, '2 Tailor and review both documents')
    steps(d, [
        'In Section 1, optionally enter the job title and company. These details help label the saved version and personalize the letter.',
        'At “Please add the job description so your resume can be tailored.” paste the full posting. Include responsibilities, tools, skills, and qualifications. Use at least 80 characters and 12 words; the maximum is 40,000 characters.',
        'Choose the employer’s Applicant tracking system if known. Otherwise leave General ATS format selected. The selector includes all 46 requested systems.',
        'Open Tailoring preferences if needed. Focused keeps up to four bullets for the first role and three for later roles. Full detail keeps every bullet. Both retain the original roles, employers, titles, and dates.',
        'Use Additional keywords for a job phrase the app missed, separating phrases with commas. The posting and your master resume must support any phrase used to improve the score.',
        'Choose Tailor resume + cover letter. Both documents appear in Section 2, and the application is added to Saved versions when browser storage is available.'
    ])
    h(d, 'What to check in Section 2', 2)
    table(d, ['Tab','Purpose'], [
        ('Resume and Cover letter','Read the complete documents. Confirm accuracy, relevance, dates, contact details, and the letter’s greeting.'),
        ('Keywords','Compare separate resume and cover-letter scores, included phrases, missing wording, and unsupported qualifications.'),
        ('Writing review','Check English edits, action verbs, sentence length, and wording flagged for your review.'),
        ('ATS review and Changes','Check document content and platform guidance, then inspect how the original resume was tailored.')
    ], [1.8,5.2])
    h(d, 'Understand the 10 out of 10 target', 2)
    para(d, 'The score measures the unique job phrases this app detects. A phrase counts only when it appears in the document and is supported by your master. Unsupported requirements keep the score below 10/10. Cover-letter scoring uses its paragraph text. Add genuine missing experience to your source file, upload it again and tailor; do not add qualifications you do not possess. This measure does not cover every requirement or predict an employer’s ATS result.')

    new_page(d, '3 Edit and download your application')
    h(d, 'Make final edits', 2)
    para(d, 'Select Resume or Cover letter, then choose the edit icon in the review toolbar. The resume editor revises the summary, competencies and accomplishment bullets. To change contact details, employers, job titles, dates, skills or education, edit your original master file and upload it again. In the cover-letter editor, personalize the company, subject, greeting, paragraphs, closing and signature.')
    para(d, 'Save your changes and recheck Keywords and Writing review. Editing a resume regenerates the cover letter until you have customized the letter. Once customized, the letter is preserved when the resume changes, so check that both documents still agree. Changes to a tailored document do not rewrite your master.')
    h(d, 'Save a resume and a cover letter', 2)
    steps(d, [
        'In the export area, select Resume or Cover letter. Choosing the corresponding preview tab also selects that document.',
        'Choose Word .docx, PDF, or Plain text. This prepares a file and opens the save options; it is not the final save step.',
        'If you are using the app inside ChatGPT, choose Open download page. A separate browser tab receives the selected document.',
        'Choose Save file… if a folder picker is available, or Download file to use the browser’s download manager. A phone or tablet may show Share file with a Save to Files option.',
        'Check the chosen folder or the browser’s Downloads list, open the file, and confirm the content. Repeat for the other document.'
    ])
    table(d, ['Format','Choose it when'], [
        ('Word DOCX','You want an editable file or the employer requests a Word resume.'),
        ('PDF','The employer accepts PDF and you want stable layout with selectable text.'),
        ('Plain text TXT','You need to copy content into a form or read text without formatting. Some ATS resume fields do not accept TXT.')
    ], [1.45,5.55])
    para(d, 'Files are generated in your browser. The app cannot force every browser to use one download folder. A native save picker lets you choose; ordinary downloads follow your browser settings. “Download requested” asks you to check the download manager. “File saved” is shown only after the native file write completes. If a file opens in a preview, use that preview’s Download or Share action.')
    para(d, 'Choose the copy control to copy the selected document. If automatic copying is unavailable, use Select all text in the copy dialog, then Control+C, Command+C, or your device’s Copy command.')

    new_page(d, '4 Manage saved data and shared links')
    para(d, 'Saved versions holds your last ten applications in the current browser profile, including each resume, cover letter, posting, and review data. Open Saved versions and choose an application to restore it into the workspace. Older saved resumes receive a cover letter when opened if one was not previously saved.')
    table(d, ['Action','What happens'], [
        ('Open or refresh the app','Sections 1 and 2 return to blank. Job fields, previews, temporary editor changes, and prepared download links reset. The saved master and versions remain.'),
        ('Open a saved version','The selected application fills the workspace. Its saved master snapshot is available with that application.'),
        ('Clear saved versions','Immediately deletes saved applications from this browser, including resumes and cover letters, and clears the open workspace. The master remains.'),
        ('Clear Master Resume','Immediately deletes all master sections and imported file data from the app. Saved applications remain, and the master stays blank when you reopen the app.'),
        ('Use a different device or browser profile','The new profile starts without your master or history. There is no account synchronization of applicant data.')
    ], [2.1,4.9])
    h(d, 'Sharing the app', 2)
    para(d, 'Share the ordinary app link. A new end user with access to the Site and a browser profile that has not used it receives a blank Master resume and zero saved versions. The link does not transfer your master or saved applications. Site access follows the owner’s sharing settings; sending the link does not itself grant access.')
    para(d, 'Without password protection, people using the same device and browser profile can see the local data stored there. On a shared computer, download the files you need, clear saved versions, and clear the master. Clearing app data does not delete files already saved to Downloads or documents transferred into a separate download tab. Close those tabs and manage downloaded files separately.')
    h(d, 'When storage is unavailable', 2)
    para(d, 'A persistent notice explains when storage is unavailable. Upload a master, tailor and download documents in the current tab, then save copies before closing it. If clearing reports a storage error, retry and verify it; the prior stored source may return after reopening. Confirmed deletion is reported only after storage accepts the change.')

    new_page(d, '5 Accessibility and common questions')
    h(d, 'Navigate with a keyboard or assistive technology', 2)
    bullets(d, [
        'Use Tab and Shift+Tab to move through controls. Use Skip to main content to reach the workspace.',
        'Within the review tabs, use Left Arrow, Right Arrow, Home, or End to switch panels. Tab then moves into the selected panel.',
        'Use Escape or the labeled Close button to close a dialog. Focus returns to the control that opened it.',
        'Use browser zoom and text settings to enlarge content. Document previews and keyword tables can receive keyboard focus for scrolling.',
        'Use Plain text exports or the manual copy dialog if these better suit your reading tools. Motion follows the device’s reduced-motion preference.'
    ])
    para(d, 'The app includes accessibility improvements and automated checks. Real screen-reader, device, and browser verification is still needed before a full accessibility conformance claim.')
    h(d, 'Resolve common problems', 2)
    table(d, ['Problem','Next action'], [
        ('Tailoring is disabled','Upload a readable source file; no resume sections or personal fields are mandatory.'),
        ('The posting is rejected','Paste the full description. Check the minimum of 80 characters and 12 words and the 40,000-character maximum.'),
        ('The score is below 10/10','Read the missing phrases and evidence gaps. Correct genuine omissions in the master and tailor again.'),
        ('A file is missing','Finish the save step, check Downloads or the selected folder, and use Open download page from an embedded app.'),
        ('Word or PDF fails to load','Try the format again. Use Plain text if the export dependency cannot load.'),
        ('PDF changes a name or character','Use Word or Plain text to preserve characters outside the PDF font’s supported set.'),
        ('My saved versions are absent','Return to the original browser profile and unlock through Privacy and password if protected. Clearing browser data or private browsing can remove or isolate saved data.')
    ], [2.1,4.9])
    para(d, 'For feedback, describe the task, device, operating system, browser version, expected result, actual result, and any error text. Remove personal resume details from screenshots or examples you send. The Web and Mobile Browser Guide contains the developer’s detailed troubleshooting and test plan.')
    save(d,'Le_Resume_Studio_User_Tutorial')

def developer():
    d = base('Developer Guide', 'This guide documents the deployed static app, its HTML and CSS structure, JavaScript modules, data contracts, browser APIs, export pipeline, tests, and publishing process. It describes the implementation as it exists, including its limits, so a developer can maintain it without assuming a backend or AI service.')
    h(d, '1 Runtime and local setup')
    table(d, ['Technology','Actual role'], [
        ('HTML5','index.html and download.html define semantic forms, previews, native dialogs, navigation, and save controls.'),
        ('CSS','styles.css, accessibility.css, and browser.css provide the visual system, reflow, focus, contrast, and mobile adjustments.'),
        ('JavaScript','Native ES modules run the entire application in the visitor’s browser. There is no transpilation or app build step.'),
        ('Vendored libraries','docx.js writes DOCX files; pdf-lib.min.js writes PDFs. Nimbus Sans font assets support the PDF layout. JSZip 3.10.1 reads DOCX; PDF.js 5.6.205 extracts PDF text. Reader assets and licenses are served locally.'),
        ('Node.js','Development tooling runs node:test checks and export validation scripts. It is not an application server.'),
        ('Python and Java','Neither is used by the app runtime. Python builds these guides with python-docx and can run a local static server or document QA. No Java backend or native Android project is present.')
    ], [1.5,5.5])
    para(d, 'The repository package is an ES-module package with a test script and no installed application dependency requirement. Use a recent Node.js release for development. Serve dist over HTTP or HTTPS; opening index.html as a file URL is not a supported module-loading workflow.')
    code(d, 'npm test\n# Optional local server from the repository root\npython3 -m http.server 8000 --directory dist\n# Open http://localhost:8000')
    para(d, 'Production serves dist through Sites using .openai/hosting.json. There is no application REST API, database, account synchronization, OpenAI API, or required API key. Network requests load static files and export assets. The app’s API consists of the JavaScript interfaces described below and browser APIs listed in Section 5.')

    new_page(d, '2 Source map and presentation layers')
    table(d, ['File or folder','Responsibility'], [
        ('dist/index.html','Main page, input/output sections, six review tabs, eight dialogs, labels, and export controls.'),
        ('dist/app.js','DOM orchestration, upload-only master, saved history, document edits, reset lifecycle, clipboard, and export dialogs.'),
        ('dist/resume.js','createBlankMaster and roleHeaderFields. Uploaded header order is retained; source headings retain role/company/years.'),
        ('dist/privacy.js','Web Crypto password-derived encrypted vault, unlock, write queue and lock invalidation.'),
        ('dist/imports.js','imports.js reads DOCX/PDF/TXT/RTF and validates facts; formatting.js detects bounded typography and section order.'),
        ('dist/engine.js','Keyword catalog and detection, evidence support, prioritization, tailoring, cover letters, text output, and scoring.'),
        ('dist/writing.js','The HR writing profile, nine action-verb categories, conservative English edits, and writing review.'),
        ('dist/ats.js','46 named systems plus general profile, shared template rules, source-linked guidance, and file checks.'),
        ('dist/exports.js','Resume and cover-letter DOCX/PDF generation, font embedding, text wrapping, and pagination.'),
        ('dist/downloads.js','Payload validation, filenames, Blob creation, private session transfer, save links, and optional file sharing.'),
        ('dist/download.html and download-page.js','Standalone download page receiving an exact document snapshot.'),
        ('dist/accessibility.js','Dialog focus, announcements, field errors, and focus restoration.'),
        ('dist/browser.js','Clone fallback, lazy export-library loading, and visible viewport tracking.'),
        ('dist/vendor','Browser export libraries, font assets, and relevant font license files.'),
        ('tests','node:test suites, DOM adapter, parsing helpers, source fixtures, and export QA scripts. Not deployed.'),
        ('README and review Markdown files','Implementation notes for ATS, accessibility, and browser compatibility.')
    ], [2.65,4.35])
    h(d, 'HTML and CSS maintenance', 2)
    para(d, 'Both pages declare English, allow viewport zoom, and load styles.css, accessibility.css and browser.css in that order. They use JavaScript modules. Preserve stylesheet order for layout, focus, contrast, reflow and mobile overrides.')

    new_page(d, '3 Data contracts and stored records')
    para(d, 'Master and resume objects are plain serializable data. The blank factory returns empty strings and arrays. A readable upload saves without required applicant fields or section headings. Imports add source, importedText, importedFileName and formatting metadata. Clear Master Resume removes the unprotected master record and keeps only a content-free deletion flag, preventing history recovery. Protected storage writes a blank master snapshot. Saved versions retain separate snapshots until cleared.')
    code(d, 'const master = {\n  name: "Avery Taylor", phone: "",\n  email: "avery@example.com", linkedin: "",\n  source: "Uploaded master resume",\n  summary: "Business analyst supporting reporting.",\n  competencies: ["SQL", "Requirements Analysis"],\n  roles: [{\n    employer: "Example Company", title: "Business Analyst",\n    dates: "2022 to Present",\n    bullets: ["Analyzed SQL reports for stakeholders."]\n  }],\n  skills: [{group: "Analysis", items: ["SQL"]}],\n  education: [{school: "Example University",\n    qualification: "Bachelor of Science", dates: "2020"}]\n};')
    para(d, 'The example is fictitious. All applicant fields are optional for upload. Shape validation permits empty strings and arrays. Unlabeled text is retained as source summary content when no structure is recognized; missing facts are not invented. Imported text makes a saved upload ready even when the name is absent. Export payloads allow a blank name. File-format, size and readable-text protections still apply.')
    table(d, ['Record','Important fields'], [
        ('Application','id, createdAt, resume, master, jd, title, company, length, extra, ats, keywords, changes, englishCorrections, writingProfile, coverLetter, audit, coverAudit, review.'),
        ('Cover letter','date, company, subject, salutation, paragraphs as strings, closing, signature; edited marks customized content.'),
        ('Keyword','Detected jobTerms, normalized concept name, evidence/support status, and inclusion information. Preserve the evidence flag when rescoring.'),
        ('Audit','total, before, after, supported, coverage, score, targetMet, targetPossible, possibleCoverage, phrases, termsNeeded, evidenceNeeded.')
    ], [1.4,5.6])
    para(d, 'Store records as JSON. cloneData uses structuredClone when available and a JSON clone fallback for these plain objects. Startup shape checks discard corrupt history entries and cap history at ten. Do not treat externally edited localStorage as trusted executable content.')

    new_page(d, 'Master resume import interfaces')
    code(d, 'validateResumeFile(file) -> extension or throws\nextractResumeText(file, {signal, loadZip, loadPDF, onFormatting} = {})\n  -> Promise<string>\nparseResumeText(text) -> parsed source object')
    para(d, 'File validation accepts PDF, DOCX, TXT and RTF with a 10 MB limit. TXT is strict UTF-8. RTF decodes Windows-1252 and supported Unicode controls, excludes metadata and embedded objects, and never executes content. DOCX uses JSZip with a 30 MB expanded archive limit. PDF.js reads up to 30 pages with evaluation disabled and local assets. No OCR is provided. Text is capped at 200,000 characters.')
    para(d, 'Parsing recognizes common sections and confirmed three-field employment headings. Word numbering preserves bullet markers; PDF text is ordered by row and horizontal position. Until a complete next heading is confirmed, text remains with the current role, including its final accomplishment. headerOrder preserves role, employer and date order. Readable uploads save automatically even with missing facts or headings. Unreadable or invalid files preserve the prior source. Roles without bullets do not produce accomplishment paragraphs in the cover letter.')
    para(d, 'app.js guards reads with an epoch and AbortController. Clear, dialog close and workspace reset invalidate pending reads; late completion cannot restore deleted content. Upload disables the picker while reading; Clear can cancel a pending read. The input clears before its chooser opens and after completion, allowing a rejected selection to be retried after renaming. Before extraction, filenames are normalized with NFKC, trimmed and compared without case against the current master filename and its uploadedFileNames array. Only successful imports reserve a name; the array is saved alongside the master through the existing browser storage or encrypted vault. Clear Master Resume discards this history. Names are never checked across users. Applicant names are not duplicate blockers. Progress, associated errors and a read-only text view support keyboard access. Clear resets the file selection, extracted text, master marker and current output; stored versions remain separate. It calls resetWorkspace immediately after blanking the master and before awaiting storeValue, so text boxes clear even while encrypted storage is still writing.')
    para(d, 'PDF.js and its worker must be served as modules with standard_fonts and licenses. Readers load only when needed. If DOCX or PDF reading fails, revise the file or upload TXT or RTF; manual master entry is unavailable. formatting.js normalizes bounded typography, margins, page size, labels and section order. Word uses detected font settings; PDF infers a compatible serif, sans or mono family. Tables, graphics and arbitrary original positions are not reproduced.')

    para(d, 'TXT and RTF use a single column, Arial 11 point and 0.75 inch margins, within MIT recommended font and margin ranges. This fallback does not guarantee ATS acceptance.')
    link(d, 'MIT Resume Checklist', 'https://capd.mit.edu/resources/resume-checklist/')
    new_page(d, '4 Tailoring engine API and score behavior')
    h(d, 'Primary JavaScript interfaces', 2)
    code(d, 'tailorResume(master, jd, {\n  title = "", company = "", length = "balanced",\n  extra = "", ats = "general"\n} = {}) -> Application')
    para(d, 'Synchronous entry point. Rejects fewer than 80 trimmed characters or 12 words, more than 40,000 characters, or an entirely empty master. balanced retains up to four first-role bullets and three later-role bullets; detailed keeps all. It clones the master, creates both documents, and returns reviews. Catch its Error and show a readable message without enabling exports.')
    code(d, 'extractKeywords(jd, master, extra = "") -> Keyword[]\nauditResume(resume, keywords, master) -> Audit\nauditCoverLetter(letter, keywords) -> Audit\nreviewApplication(result, {reextract = false} = {})\n  -> result with refreshed audits and writing review')
    para(d, 'Detection combines a finite catalog with named-tool and stated-skill detection. Additional comma-separated terms help recover missed phrases. Catalog concepts can have multiple employer aliases; every unique normalized job phrase is scored. Additional terms absent from the posting are excluded from the denominator. Resume evidence controls support; typing an unsupported term into the output cannot create evidence.')
    code(d, 'createCoverLetter(application) -> CoverLetter\nresumeText(resume) -> string\ncoverLetterText(letter, resume) -> string')
    para(d, 'The letter draws on supported accomplishments and job terms. Text serializers provide the source for TXT exports and clipboard content. Cover scoring examines paragraphs rather than the subject or contact metadata. reviewApplication mutates the review fields on the passed application; it does not replace a customized letter.')
    h(d, 'Pipeline and invariants', 2)
    para(d, 'The engine detects phrases, checks evidence, focuses the summary, moves relevant competencies first, ranks bullets within each role, rewrites equivalent supported wording, orders skills, inserts missing supported phrases, polishes English, creates the letter, and audits both documents. Original employers, titles, dates, metrics, and qualifications are preserved. Support work must retain its responsibility level.')
    para(d, 'The 10/10 target requires a nonzero denominator and every detected phrase to be both supported and included. Percent coverage and the 0-to-10 score are floored to one decimal, so an incomplete match cannot round up to 10. Unsupported phrases stay in the denominator; zero detected phrases does not count as success. The catalog is finite, and the overlap measure is not a live ATS score or semantic evaluation of every job requirement.')
    para(d, 'writing.js exposes polishResume, polishCoverLetter, cleanEnglish, polishBullet, actionVerb, firstPersonBullet, and writingReview. The HR profile models a hiring reviewer with over twenty years of experience as an editorial lens. It neither adds that tenure to the applicant nor performs human review. Grammar rules are conservative; unrecognized openings and long sentences remain review items.')

    new_page(d, '5 Export interfaces and browser APIs')
    code(d, 'createDownloadPayload(application, kind, format) -> Payload\nvalidateDownloadPayload(payload) -> payload or throws\ncreateExportBlob(payload, dependencies = {}) -> Promise<Blob>\nencodeDownloadPayload(payload) -> base64url string\ndecodeDownloadPayload(hash) -> validated Payload\ndownloadPageUrl(payload, host) -> private URL string\nconfigurePrivateHandoff(link, payload, status, host) -> cleanup\nreadDownloadTransfer(host = globalThis) -> Payload')
    para(d, 'kind is resume or cover; format is docx, pdf, or txt. The version-1 payload contains name, resume, letter, and ats. Resume payloads contain the selected resume; cover payloads contain contact data and the letter. They exclude the full master, job posting, and saved history. Filenames combine applicant, job title, company, and document type, sanitize punctuation, and add the extension.')
    para(d, 'Validation constrains supported formats, filename characters and extension, contact strings, arrays, and letter paragraphs. configurePrivateHandoff opens a same-origin tab on the user’s click. A 600,000-byte maximum snapshot is stored in sessionStorage under a random UUID for five minutes; only #transfer=UUID enters the URL. The opened tab inherits a copy of the source sessionStorage, consumes the transfer, clears its URL and retains the snapshot for refresh. The parent removes its transfer record. A copied URL in another session fails. Blocked popups or sessionStorage must offer direct save. Legacy encode/decode helpers remain internal serialization utilities and are not used in public links.')
    code(d, 'createDocx(resume, docxApi) -> Promise<Blob>\ncreateCoverDocx(letter, resume, docxApi) -> Promise<Blob>\ncreatePdf(resume, pdfApi, fontPrograms) -> Promise<Uint8Array>\ncreateCoverPdf(letter, resume, pdfApi, fontPrograms)\n  -> Promise<Uint8Array>')
    para(d, 'createExportBlob wraps PDF bytes in the proper Blob type and loads the requested local library. TXT uses no export library. Source formatting travels with resume payloads and snapshots; cover letters keep their professional template. PDFs embed matching CFF font programs: Nimbus Sans or PDF.js serif/mono assets. Non-ASCII normalization may alter characters, so the review recommends Word or TXT when needed.')
    table(d, ['Browser API','Purpose and fallback'], [
        ('localStorage and sessionStorage','Browser-owned master/history and temporary download snapshot. Catch blocked access or quota errors.'),
        ('Blob and URL object URLs','Prepare downloadable bytes. Revoke old links when replaced, cleared, or reset.'),
        ('showSaveFilePicker','Optional save picker outside frames; confirms completion after stream.close. Otherwise use a download link [1].'),
        ('navigator.share and canShare','Optional exact-file sharing on a direct click; hidden if unsupported or embedded [2].'),
        ('navigator.clipboard.writeText','Automatic copy when permitted; otherwise show selectable text [3].'),
        ('HTMLDialogElement and VisualViewport','Modal focus and mobile keyboard sizing. CSS remains the viewport fallback [4,5].')
    ], [2.3,4.7])

    new_page(d, '6 HTML CSS and editor implementation')
    h(d, 'Preserve semantic relationships', 2)
    para(d, 'index.html uses labeled forms, a skip link, navigation and main landmarks, alert regions for validation, and status regions for asynchronous updates. The review tablist contains Resume, Cover letter, Keywords, Writing review, ATS review, and Changes. Each tab controls a named panel. app.js keeps one tab selected, uses a roving tabindex, and handles Left/Right/Home/End. Preserve these relationships when adding or renaming UI elements [6].')
    code(d, '<label for="job-description">Please add the job description\n  so your resume can be tailored.</label>\n<textarea id="job-description" required minlength="80"\n  maxlength="40000"\n  aria-describedby="jd-helper jd-requirements char-count">\n</textarea>')
    para(d, 'Custom JavaScript also validates the word count because the form uses novalidate. Editors identify the invalid field, associate its error, expand a closed role if necessary, and focus it. Named dialog headings receive initial focus; closing restores focus to the opener. Status messages are announced inside an active modal instead of its inert background [7].')
    h(d, 'Upload and clear controls', 2)
    para(d, 'The master dialog restores version 1 section grouping with read-only contact, summary, competencies, a separate expanded section for every detected role, technical skills and education. showImportedText renders escaped facts and bullets, plus the original-text disclosure. The native picker remains the only source input. There is no master form, manual Save or Add/Remove handler. buildEditor and collectEditor are used only for tailored output. Source corrections happen in the original file and are uploaded again. Clear cancels pending reads and resets all displayed data before awaiting storage. keepMasterOpen leaves the empty dialog visible. deleteStoredMaster removes the master key after saving a boolean deletion flag; encrypted storage writes a blank snapshot. The direct click avoids native confirm dialogs that embedded browsers may suppress. masterDeletionFailed enables retry after a failed write. Storage events invalidate a cleared master in other tabs.')
    h(d, 'Responsive CSS and focus', 2)
    para(d, 'The base layout uses a sidebar and a two-column workspace. Overrides use minmax(0, ...), min-width:0, wrapping controls, relative type sizes, and narrow-screen stacking. Browser rules switch layouts around 1000 and 640 CSS pixels, add safe-area insets, and account for short landscape screens. Keep form text at least 16 CSS pixels at default size and preserve browser zoom.')
    code(d, '/* Representative rule from browser.css */\n@media (max-width:640px) {\n  .export-actions {\n    display:grid;\n    grid-template-columns:repeat(2,minmax(0,1fr));\n  }\n}')
    para(d, 'Accessibility rules enlarge targets, show clear focus indicators, support forced colors and reduced motion, and make scrollable previews keyboard reachable. Viewport tracking updates --visible-height and --visible-top when the mobile keyboard changes the visible area; pinch zoom removes overrides. These measures require real-device and assistive-technology verification before a conformance claim.')

    new_page(d, '7 State lifecycle privacy and failure handling')
    table(d, ['Storage key','Contents and lifetime'], [
        ('lw-resume-studio-master-v1','Saved master in localStorage for this origin and browser profile. Clear removes this record, including imported text, filename and formatting.'),
        ('lw-resume-studio-history-v1','Up to ten complete application records. Each includes a master snapshot and both documents.'),
        ('lw-resume-studio-export-v1','Selected document snapshot in the standalone download tab’s sessionStorage.'),
        ('lw-resume-studio-draft-v1','Retired. Removed at startup; job drafts are no longer persisted.')
    ], [3.0,4.0])
    h(d, 'Startup and resets', 2)
    para(d, 'When protection is off, load a valid saved master or recover this browser’s history snapshot. The boolean lw-resume-studio-master-cleared-v1 flag prevents recovery after clearing. Legacy blank markers remain supported. A protected workspace starts blank until unlocked. Master readiness controls tailoring; new browser profiles receive no applicant data.')
    para(d, 'resetWorkspace runs at startup and on pageshow, including back/forward-cache restoration. It clears job fields, output, audits, dialogs, manual-copy text, preferences, and prepared download links, while retaining master/history. workspaceEpoch guards asynchronous tailoring and export completion so a reset or deletion cannot repopulate cleared output. Preserve these guards when adding new asynchronous work.')
    h(d, 'Deletion and storage errors', 2)
    para(d, 'Clear saved versions deletes the history key when unprotected, or awaits an encrypted empty-history write when protected. It clears memory, output, editors and prepared links, cancels pending work and reports persistence failures. Clear Master Resume deletes the source record, or saves an encrypted blank master, and keeps application history. Other-tab storage changes invalidate stale copies.')
    para(d, 'If storage is blocked or full, reads and writes are caught, a persistent notice appears, and the current tab can continue tailoring and downloading. Saving or clearing a master reports whether persistence failed. The user must download copies before the tab closes. Existing downloaded files and separate download-tab snapshots are outside history deletion.')
    h(d, 'Security and sharing boundaries', 2)
    para(d, 'Applicant data stays local. Opt-in password protection encrypts saved records and removes legacy plaintext. Without it, users sharing a browser profile can read saved data. Normal app and download URLs contain no applicant content. Site sharing exposes app code; no applicant database or cloud account exists.')
    para(d, 'Escape stored text, validate transfers and avoid executable interpolation. Both HTML pages use no-referrer metadata and same-origin script/connection policies, block object embeds and external form submission, and load local vendors. No service worker or native offline package is shipped.')

    new_page(d, 'Password protected storage API')
    code(d, 'PrivateVault(storage, crypto)\n  exists() -> boolean\n  create(passphrase, {master, history}) -> Promise<void>\n  unlock(passphrase) -> Promise<{master, history}>\n  write({master, history}) -> Promise<void>\n  flush() -> Promise<void>\n  lock() -> void')
    para(d, 'privacy.js uses Web Crypto AES-GCM with a 256-bit nonextractable key, a random 16-byte salt and a fresh 12-byte IV for each write. PBKDF2-HMAC-SHA-256 derives the key with 600,000 iterations. The minimum passphrase length is 12 characters. localStorage key le-resume-studio-private-vault-v1 contains schema, salt, IV, work factor and authenticated ciphertext; it contains no plaintext master, history, password or key. Authentication failure rejects wrong passwords and modified ciphertext.')
    para(d, 'Writes serialize through a promise queue and capture a JSON snapshot. Lock increments an epoch and drops the key; a pending encryption cannot write after locking. Encrypted writes replace the vault atomically only after encryption succeeds. Upload and deletion controls await persistence. Other-tab vault updates invalidate an unlocked copy. Storage failure preserves the current tab and exposes an error rather than claiming durable protection. Protection setup and unlock both retry deletion of legacy plaintext keys.')
    para(d, 'This is local passphrase protection, not a server login or a password recovery system. It does not hide an unlocked tab, protect downloads on the filesystem, secure a compromised device, or prevent a trusted future app script from reading an unlocked workspace. Browser session restoration and actual popup/iframe behavior require real-device verification. Close download tabs when handing over a shared device. Keep the app origin and vendored assets trusted.')
    link(d, 'OWASP password storage guidance', 'https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html')
    link(d, 'MDN sessionStorage and opener inheritance', 'https://developer.mozilla.org/en-US/docs/Web/API/Window/sessionStorage')

    new_page(d, '8 Verification deployment and maintenance')
    h(d, 'Checks and release procedure', 2)
    para(d, 'The master review and deletion release passed 133 node:test checks on October 14 2026. Tests cover tailoring evidence, both document scores, English rules, actual export bytes, transfers, save failures, history deletion, reset guards, keyboard semantics, capability fallbacks, and seven fresh-master journeys. DOM adapters and capability simulations do not substitute for browser rendering or real ATS uploads.')
    steps(d, [
        'Run npm test. If export code changes, also run the applicable export-check scripts and inspect actual Word/PDF output, extracted text, page breaks, and fonts.',
        'Confirm dist contains only deployable assets. Keep tests and personal fixtures out of the hosting archive. Review README, ATS_COMPATIBILITY, ACCESSIBILITY, and BROWSER_COMPATIBILITY when behavior changes.',
        'Use the Sites source workflow to obtain a temporary repository write credential, open the checkout, commit and push changes, and package the static dist directory.',
        'Save the Site version with its project ID, exact commit SHA, and archive, then deploy that saved version. Preserve the existing sharing audience. Inspect deployment status and keep the previous saved version available for rollback.'
    ])
    para(d, 'The Sites management API is a publishing interface, separate from the browser app: sites_create_source_repository_write_credential opens authorized source access; sites_save_site_version records commit_sha and archive; sites_deploy_site_version publishes the version_id. Keep credentials out of source files, URLs, logs, and client JavaScript. The hosting manifest uses static.directory = dist.')
    h(d, 'Safe extension points', 2)
    bullets(d, [
        'Add job vocabulary in engine.js with explicit evidence and phrase-boundary tests. Test alternative terms and unsupported qualifications; do not lower the score denominator to meet a target.',
        'Extend writing.js conservatively. Test tense, English punctuation, metrics, and responsibility level against the source.',
        'Update ATS profiles only with current primary vendor guidance. Most selected systems use general guidance; this app has no live ATS integration or certification.',
        'For a future AI service, synchronized accounts, or native app, design a separate authenticated backend, data handling, and mobile packaging plan. None of these components exists in the present app.'
    ])
    h(d, 'Primary API and accessibility references', 2)
    refs = [
        ('[1] MDN showSaveFilePicker','https://developer.mozilla.org/en-US/docs/Web/API/Window/showSaveFilePicker'),
        ('[2] MDN Navigator share','https://developer.mozilla.org/en-US/docs/Web/API/Navigator/share'),
        ('[3] MDN Clipboard writeText','https://developer.mozilla.org/en-US/docs/Web/API/Clipboard/writeText'),
        ('[4] MDN HTML dialog','https://developer.mozilla.org/en-US/docs/Web/HTML/Reference/Elements/dialog'),
        ('[5] MDN VisualViewport','https://developer.mozilla.org/en-US/docs/Web/API/VisualViewport'),
        ('[6] W3C tabs pattern','https://www.w3.org/WAI/ARIA/apg/patterns/tabs/'),
        ('[7] W3C modal dialog pattern','https://www.w3.org/WAI/ARIA/apg/patterns/dialog-modal/')
    ]
    for label,url in refs:
        p=link(d,label,url); p.paragraph_format.space_after=Pt(3)
    save(d,'Le_Resume_Studio_Developer_Guide')

def browsers():
    d = base('Web and Mobile Browser Guide', 'Use this guide to investigate browser and device differences, reproduce failures, and verify the app before making compatibility claims. The app includes capability fallbacks and responsive layouts. Real Safari, Chrome, mobile, screen-reader, and live ATS behavior has not been certified by the automated checks.')
    h(d, '1 Support model and current limits')
    para(d, 'Intended targets are current Safari on macOS/iPhone/iPad, Chrome on desktop/Android, and desktop Edge and Firefox. The app requires JavaScript modules and native HTML dialogs. It uses capability detection for optional APIs rather than promising identical controls on every browser. It is a website, with no native iOS/Android package or service worker.')
    table(d, ['Capability','Implemented behavior','Status'], [
        ('Folder picker','Save file… when showSaveFilePicker is available outside an embedded frame; otherwise Download file.','Simulated'),
        ('File sharing','Share file only when File, share, and canShare accept the exact file, outside frames.','Simulated'),
        ('Clipboard','Direct writeText on click; manual selectable-text dialog after rejection or absence.','Simulated'),
        ('Browser storage','Persistent notice and current-tab operation after blocked reads/writes or full storage.','Simulated'),
        ('Mobile keyboard','VisualViewport height/offset update; CSS viewport fallback; overrides removed during pinch zoom.','Simulated'),
        ('Clone and syntax','JSON clone fallback; application avoids literal regex lookbehind, Array.at, and Object.hasOwn.','Simulated'),
        ('Layout and accessibility','Responsive CSS, safe areas, zoom, focus, tab semantics, and reduced motion.','Code checks'),
        ('Actual device interactions','Download managers, share sheets, rendering, rotation, keyboards, and screen readers.','Needs testing')
    ], [1.45,4.35,1.2])
    para(d, 'Simulated means the feature or failure branch was exercised with test adapters. It does not certify a particular browser version. A supported optional API can still be denied by permissions, security policy, an embedded context, or user cancellation. Record the browser and operating system when verifying a release.')

    new_page(d, '2 Download and export troubleshooting')
    h(d, 'A prepared file does not appear in Downloads', 2)
    para(d, 'Reproduce by choosing a document and format, then inspect the save dialog. Format buttons prepare the Blob; the user must choose Save file… or Download file. In an embedded app, use Open download page. Ordinary downloads follow browser settings and may open in a preview. The app cannot set a universal destination folder. MDN documents the picker as limited in availability, requiring a secure context and user activation [1].')
    para(d, 'Inspect downloads.js configureSaveLink and canChooseLocation, plus download-page.js. Native picker success is reported only after write and close. An AbortError means cancellation; another save error switches the retained link to a normal download. Keep the Blob link alive until replacement/reset and do not automatically click then revoke it. Check browser download permissions and the preview’s own Download action.')
    h(d, 'Sharing is missing or fails', 2)
    para(d, 'Absence of Share file is expected if exact-file sharing is unsupported or the app is embedded. configureShareButton checks canShare before displaying it and invokes share directly from a user click. A permissions policy or loss of transient activation can reject sharing [2]. Cancellation leaves the download available. Promise resolution is not treated as proof of a file saved to the device.')
    new_page(d, 'Master resume upload troubleshooting')
    para(d, 'Upload DOCX, text-based PDF, UTF-8 TXT or RTF. Limits are 10 MB, 30 PDF pages and 200,000 extracted characters. Convert older .doc files, remove PDF passwords and use text copies of scans. Uploads save automatically without required sections, names, skills, dates or job headings. The master uses section grouping like version 1, with a separate expanded section for every role; details remain read-only. Correct the original file and reupload if parsing misses facts; each successful upload must use a new filename, ignoring capitalization. Duplicate errors preserve the current master; failed uploads do not reserve names. Clear Master Resume resets filename history. Clear Master Resume immediately empties the uploaded-text, job-description, extra-keyword and copy-text boxes before deleting the saved master record. It shows No master resume saved, cancels pending imports and stays blank after reopening. No native confirmation popup is used. Failed deletion remains retryable; other tabs clear their stale master. The app does not delete your source file on disk. Saved versions remain separate.')
    para(d, 'For PDF failures, check local PDF.js modules, standard fonts, MIME types and worker policy. For DOCX, check JSZip. Readers do not upload resume contents. A cloud-picked file must be downloadable by the browser. If a reader fails, use TXT or RTF. Word and PDF styling is detected conservatively; missing fonts use compatible replacements. Native Safari, Chrome, mobile and assistive-technology upload interactions still require real-device checks.')
    h(d, 'Word or PDF export fails while Plain text works', 2)
    para(d, 'Check browser.js loadExportLibrary and network responses for vendor/docx.js or vendor/pdf-lib.min.js. Libraries are loaded on demand, concurrent requests are deduplicated, and a failed load can be retried. PDF generation also needs the vendored font programs. Serve correct JavaScript MIME types over HTTP/HTTPS and check for missing files, policy blocks, or rejected fetches. Plain text has no export-library dependency.')
    h(d, 'The PDF changes characters or the ATS rejects a file', 2)
    para(d, 'exports.js normalizes PDF text into a limited Latin character set; characters outside it may become question marks and accents may be removed. ats.js reviews this and warns users to use Word or TXT to preserve text. An employer’s upload restrictions remain authoritative. Choose a supported file type and check the actual prepared size; the app’s profiles include source-backed limits for Greenhouse and Workable. After upload, inspect parsed fields. No live parser acceptance is guaranteed.')
    para(d, 'Acceptance check: prepare and open all six resume/cover-letter combinations, confirm filenames and content, cancel a picker, retry a failed save, and verify that each file exists in the chosen destination. Test both an ordinary browser tab and the embedded handoff.')

    new_page(d, '3 Storage resets transfers and copying')
    table(d, ['Symptom','Cause or expected behavior','Developer action'], [
        ('Sections 1 and 2 become blank','Expected on load, reload, restart, or pageshow. Master and history stay local.','Check resetWorkspace and workspaceEpoch guards.'),
        ('A new user sees someone’s master','Same browser profile or unintended published seed. A genuinely fresh profile must be empty.','Check local keys, createBlankMaster, and the deployment archive.'),
        ('Saved versions are absent','Different browser/profile/origin, cleared site data, private session, or storage rejection.','Check the persistent notice and caught storage errors.'),
        ('History returns after deletion','Failed persistent removal or stale cross-tab state.','Check removeItem fallback, retry status, and storage listener.'),
        ('A download page says to return','Missing/invalid fragment and no valid session snapshot, or transfer exceeds the cap.','Prepare a fresh transfer; check sessionStorage inheritance, popup permissions, expiry and validation.'),
        ('Copying is denied','Clipboard API is absent or denied. Automatic write needs an allowed secure context [3].','Use the focus-managed manual copy dialog; preserve the document text.')
    ], [1.55,3.0,2.45])
    h(d, 'Browser data is isolated by context', 2)
    para(d, 'localStorage belongs to an origin and browser profile. Private browsing has different persistence behavior, and access can throw SecurityError [4]. An embedded app may face additional storage restrictions. Treat storage as optional: the app must remain usable with its persistent warning, without claiming changes were saved when a write failed.')
    h(d, 'Transfer snapshots have a separate lifetime', 2)
    para(d, 'The standalone download page receives its selected snapshot from an inherited sessionStorage copy, using a one-use random token with a five-minute expiry. The URL contains no applicant content. The page removes the token and retains the snapshot in its own sessionStorage. Another browser session cannot retrieve a copied token. Popups and blocked session storage can prevent handoff; offer direct download instead. Clearing Saved versions resets app history and prepared links but does not erase a separate download tab’s snapshot or files already downloaded. Avoid sharing document transfer links as if they were the ordinary app link.')
    para(d, 'Acceptance check: save an application, reload and confirm blank Sections 1 and 2, explicitly reopen history, clear it, and confirm zero after reopening. Test missing master recovery from local history, explicit master clearing, a second fresh browser profile, blocked storage, clipboard denial, and deletion in another open app tab.')

    new_page(d, 'Password protection and private transfers')
    para(d, 'Privacy and password encrypts local saved data after the user chooses a passphrase. This requires HTTPS, Web Crypto and writable localStorage. A protected workspace starts blank and locked on a new load. Unlock requires the same passphrase; there is no cloud password reset. Check for denied storage, quota exhaustion and interrupted legacy-record deletion before declaring setup complete. Wrong passwords or modified ciphertext must never reveal data.')
    para(d, 'Lock must cancel stale imports and clear the workspace, editors, previews and prepared links while retaining the encrypted vault. Verify saved-version and master deletion by unlocking again. An unlocked download tab and files saved on the device remain separate copies; close or delete them separately. Different users should use separate browser profiles, or lock the protected workspace before sharing a device.')
    para(d, 'Private download handoff requires a new tab to inherit an independent copy of the opener’s sessionStorage. The URL contains only a random token. The parent removes its transfer record; the receiver consumes it within five minutes. A copied token in an unrelated session or a replay must fail. Check direct browser and embedded ChatGPT contexts, blocked popups, private browsing and Safari/mobile storage behavior. Direct save remains the fallback. No native-browser certification is claimed.')
    para(d, 'Both pages declare same-origin script and connection restrictions and no-referrer metadata. Keep vendors local and avoid adding analytics, upload endpoints or external scripts that could receive applicant content. Check browser consoles for blocked vendor behavior before extending the policy. Policy and real-device enforcement must be reviewed after deployment changes.')

    new_page(d, '4 Mobile layout focus and text issues')
    h(d, 'The keyboard hides a field or save control', 2)
    para(d, 'A phone keyboard changes the visible viewport without necessarily changing the layout viewport [5]. browser.js trackViewport listens for visible viewport resize/scroll and supplies --visible-height and --visible-top to browser.css. Dialogs use these variables with viewport fallbacks and scrolling. Check event updates while typing, rotating, and scrolling. At pinch-zoom scale other than one, overrides are removed so the browser manages zoom.')
    h(d, 'Controls overlap a notch or the screen edge', 2)
    para(d, 'Both pages include viewport-fit=cover. browser.css adds env(safe-area-inset-*) spacing for navigation, workspace, dialogs, footers, and the download page. WebKit describes these insets for edge-to-edge layouts [6]. Test portrait and landscape, including short-height landscape navigation. Keep max()/calc() fallbacks and verify computed styles on the affected device.')
    h(d, 'Text zoom causes clipping or unwanted input zoom', 2)
    para(d, 'Check stylesheet order and form font size. accessibility.css uses relative text sizing and controls at least 16 CSS pixels by default. Containers use min-width:0, wrap long labels and filenames, and change grid layouts on narrow screens. Do not fix clipping by disabling zoom. Verify 200% and 400% browser zoom or equivalent large-text settings, plus a 320 CSS-pixel viewport. Document previews and keyword tables may scroll; task controls must remain reachable.')
    h(d, 'Dialog focus or announcements behave incorrectly', 2)
    para(d, 'Inspect accessibility.js and matching HTML IDs. Each dialog needs a named, focusable heading, a Close control, and status inside the modal. Errors need field association and focus; closing restores the opener. Native modal dialogs manage the interaction boundary, but test Tab/Shift+Tab containment, Escape, background interaction, and returned focus with real assistive technology. Tabs need one active item, correct controls/labels, and arrow-key behavior.')
    h(d, 'An older browser cannot start the app', 2)
    para(d, 'The shipped implementation requires ES modules and native HTML dialogs. JavaScript-disabled users receive a notice. Existing fallbacks address optional APIs, not every old browser. Update the browser or explicitly scope and implement an accessible dialog/module fallback for a new support target. Do not claim all-browser compatibility based only on capability simulations.')
    para(d, 'Acceptance check: use a real phone keyboard in tailored editors, upload and clear a master, rotate the device, zoom, inspect long text and filenames, open dialogs and use review tabs by keyboard and screen reader. Test VoiceOver on Safari, TalkBack on Android and a desktop screen reader with supported browsers.')

    new_page(d, '5 Device verification and issue reporting')
    para(d, 'Automated evidence for this release is 133 passing tests, including upload, unique filename rejection and reset, clear and real Word/PDF extraction checks. Actual device rendering, native save/share experiences, assistive technology, and uploads into the listed ATS systems remain unverified. Record results before advertising complete compatibility or accessibility conformance.')
    table(d, ['Test environment','Priority journeys'], [
        ('Safari on macOS','Fresh master, complete tailoring/editing, all formats, clipboard fallback, zoom, keyboard and VoiceOver.'),
        ('Safari on iPhone and iPad','Keyboard viewport, rotation, safe areas, large text, downloads/previews, Share file if offered, VoiceOver.'),
        ('Chrome on Android','Same document flow, Downloads, Share file if offered, keyboard/orientation, large text, TalkBack.'),
        ('Chrome Edge Firefox desktop','All document formats, picker/download branches where available, keyboard, storage failures, screen-reader checks.'),
        ('Embedded and restricted contexts','Exact-document handoff, denied storage/copy/share, reset during generation/export, cancellation and retries.')
    ], [2.15,4.85])
    h(d, 'Release acceptance sequence', 2)
    steps(d, [
        'Start with an empty profile; verify a blank Master resume and zero versions. Add your own test master, including multiple roles, skills, and education.',
        'Tailor a supported posting and a posting with an unsupported requirement. Verify separate scores and truthful gaps, and inspect all six review tabs.',
        'Edit both documents, export all formats, inspect the saved files, and test copying, handoff, sharing, cancellation, and retry.',
        'Reload, reopen a saved application, clear history, clear the master, and verify browser-local isolation with another profile.',
        'Repeat at narrow widths, with rotation, keyboard, large text, reduced motion, forced colors where available, and assistive technology.'
    ])
    para(d, 'A useful issue report includes release version, device, OS and browser versions, embedded/direct context, storage mode, reproduction steps, expected and actual behavior, exact error, format/file size, and a sanitized example. Add severity, the affected code path, workaround, and the verified fix. Remove personal information from logs, screenshots, links, and fixtures shared with others.')
    h(d, 'Primary references reviewed October 14 2026', 2)
    for label,url in [
        ('[1] MDN showSaveFilePicker','https://developer.mozilla.org/en-US/docs/Web/API/Window/showSaveFilePicker'),
        ('[2] MDN Navigator share','https://developer.mozilla.org/en-US/docs/Web/API/Navigator/share'),
        ('[3] MDN Clipboard writeText','https://developer.mozilla.org/en-US/docs/Web/API/Clipboard/writeText'),
        ('[4] MDN localStorage','https://developer.mozilla.org/en-US/docs/Web/API/Window/localStorage'),
        ('[5] MDN VisualViewport','https://developer.mozilla.org/en-US/docs/Web/API/VisualViewport'),
        ('[6] WebKit safe areas for iPhone','https://webkit.org/blog/7929/designing-websites-for-iphone-x/')
    ]:
        p=link(d,label,url); p.paragraph_format.space_after=Pt(3)
    save(d,'Le_Resume_Studio_Web_Mobile_Browser_Guide')

tutorial()
developer()
browsers()
